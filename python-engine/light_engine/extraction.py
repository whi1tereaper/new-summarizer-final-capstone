"""Conservative text preparation and locally validated document extraction."""

from __future__ import annotations

from dataclasses import dataclass
import os
from pathlib import Path
import re
from typing import Any
from zipfile import BadZipFile

from .contract import ContractError, MAX_TEXT_CHARACTERS, MAX_UPLOAD_BYTES

_HEADING = re.compile(
    r"^\s*(?:(?:\d+(?:\.\d+)*[.)]?)\s+)?"
    r"(?:abstract|introduction|background|methods?|methodology|materials|"
    r"results?|findings?|discussion|conclusion|limitations?|recommendations?|"
    r"overview|summary|risks?|approach|architecture|implementation|rollout|"
    r"appendix|references|executive summary|current problems|current supply|"
    r"demand forecast|drought risk|options evaluated|uncertainties|idempotency|"
    r"proposed architecture|business impact|financial results?|key findings|"
    r"learning objectives|key concepts|glossary|examples|configuration|"
    r"requirements|troubleshooting|opinion|commentary|editorial|data collection|"
    r"participants|study design|data analysis|case study|market overview|"
    r"financial outlook|next steps)\s*:?\s*$",
    re.IGNORECASE,
)
_SENTENCE_START_REFERENCE = re.compile(
    r"^(?:it|they|this|these|those|such|however|therefore|moreover|thus)\b",
    re.IGNORECASE,
)
_TOKEN = re.compile(r"\b[\w'-]+\b", re.UNICODE)
_ABBREVIATIONS = re.compile(
    r"\b(?:Mr|Mrs|Ms|Dr|Prof|Sr|Jr|St|Mt|No|Fig|Eq|Ref|Sec|Dept|Univ|Inc|Ltd|"
    r"Corp|vs|etc|e\.g|i\.e|al)\.\s*$",
    re.IGNORECASE,
)
_PLACE_ABBREVIATION_END = re.compile(
    r"\b(?:[NS]\.\s*Dak|N\.?\s*[HDJMYC]|S\.?\s*[CD]|R\.?\s*I|W\.?\s*Va|"
    r"Ala|Ariz|Ark|Calif|Colo|Conn|Del|Fla|Ga|Ill|Ind|Kan|Ky|La|Md|Mass|"
    r"Mich|Minn|Miss|Mo|Mont|Neb|Nev|Okla|Ore|Pa|Tenn|Vt|Va|Wash|Wis|Wyo)\.\s*$",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class SourceSentence:
    index: int
    text: str
    start: int
    end: int
    paragraph_index: int
    section: str
    page: int | None = None


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    source_type: str
    title: str
    warnings: tuple[str, ...] = ()


def _allowed_upload_roots() -> tuple[Path, ...]:
    root = Path(__file__).resolve().parents[2]
    upload_directory = os.environ.get("SUMMARIZER_UPLOADS_DIR", "").strip()
    temporary_directory = os.environ.get("SUMMARIZER_TMP_DIR", "").strip()
    paths = [
        Path(upload_directory) if upload_directory else root / "storage" / "uploads",
        *([Path(temporary_directory)] if temporary_directory else []),
    ]
    resolved: list[Path] = []
    for path in paths:
        try:
            resolved.append(path.resolve(strict=True))
        except OSError:
            continue
    return tuple(resolved)


def _validate_upload_path(file_path: str) -> Path:
    candidate = Path(file_path)
    try:
        resolved = candidate.resolve(strict=True)
    except OSError as exc:
        raise ContractError("source_unavailable", "The uploaded document is unavailable.") from exc
    if not resolved.is_file() or not any(
        resolved != root and root in resolved.parents for root in _allowed_upload_roots()
    ):
        raise ContractError("source_path_rejected", "The document path is outside the upload directory.")
    if resolved.stat().st_size > MAX_UPLOAD_BYTES:
        raise ContractError("FILE_TOO_LARGE", "The uploaded document exceeds the supported size limit.")
    return resolved


def _extract_pdf(path: Path) -> tuple[str, tuple[str, ...]]:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError, FileNotDecryptedError

        reader = PdfReader(str(path), strict=True)
        if reader.is_encrypted:
            raise ContractError("encrypted_document", "Password-protected PDFs are not supported.")
        if len(reader.pages) > 1000:
            raise ContractError("FILE_PAGE_LIMIT", "The PDF exceeds the supported page limit.")
        pages: list[str] = []
        for page in reader.pages:
            if page.get_contents() is None:
                page_text = ""
            else:
                try:
                    page_text = page.extract_text(extraction_mode="layout") or ""
                except TypeError:
                    page_text = page.extract_text() or ""
            pages.append(page_text)
            if sum(map(len, pages)) > MAX_TEXT_CHARACTERS:
                raise ContractError("input_too_large", "The extracted document exceeds the supported text limit.")
        warnings = (
            "PDF reading order may be imperfect, especially for multi-column pages and tables.",
        )
        if not any(part.strip() for part in pages):
            raise ContractError("unreadable_document", "No readable text was found in the PDF.")
        return "\n\n".join(pages), warnings
    except ContractError:
        raise
    except ImportError as exc:
        raise ContractError("extraction_failed", "PDF extraction dependencies are unavailable.") from exc
    except Exception as exc:
        # Only parser-confirmed malformed data is described as damaged.
        from pypdf.errors import PdfReadError, FileNotDecryptedError
        if isinstance(exc, FileNotDecryptedError):
            raise ContractError("FILE_ENCRYPTED", "PDF encryption prevented extraction.") from exc
        if isinstance(exc, PdfReadError):
            raise ContractError("FILE_CORRUPTED", "PDF parser rejected malformed data.") from exc
        raise ContractError("extraction_failed", "The PDF could not be extracted safely.") from exc


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(str(path))
        blocks: list[str] = []
        if hasattr(document, "iter_inner_content"):
            for item in document.iter_inner_content():
                if isinstance(item, Paragraph):
                    value = item.text.strip()
                    if value:
                        blocks.append(value)
                elif isinstance(item, Table):
                    for row in item.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            blocks.append(" | ".join(cells))
        else:
            blocks.extend(p.text.strip() for p in document.paragraphs if p.text.strip())
            for table in document.tables:
                blocks.extend(
                    " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    for row in table.rows
                )
        text = "\n\n".join(blocks)
        if not text.strip():
            raise ContractError("unreadable_document", "No readable text was found in the document.")
        return text
    except ContractError:
        raise
    except ImportError as exc:
        raise ContractError("extraction_failed", "DOCX extraction dependencies are unavailable.") from exc
    except Exception as exc:
        from docx.opc.exceptions import PackageNotFoundError
        from lxml.etree import XMLSyntaxError
        if isinstance(exc, (BadZipFile, PackageNotFoundError, XMLSyntaxError)):
            raise ContractError("FILE_CORRUPTED", "The DOCX package is damaged.") from exc
        raise ContractError("extraction_failed", "The DOCX document could not be extracted safely.") from exc


def _looks_like_heading(line: str) -> bool:
    if _HEADING.fullmatch(line) or (
        len(line) <= 80 and line.isupper() and len(_TOKEN.findall(line)) <= 10
    ):
        return True
    tokens = line.split()
    return (
        1 < len(tokens) <= 10
        and len(line) <= 100
        and not line.endswith((".", "!", "?"))
        and all(token[:1].isupper() for token in tokens if token[:1].isalpha())
    )


def _normalize_extracted_text(text: str) -> str:
    value = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00ad", "")
    value = re.sub(r"(?<=\w)-\n(?=\w)", "", value)
    paragraphs: list[str] = []
    for raw_paragraph in re.split(r"\n\s*\n+", value):
        body_lines: list[str] = []

        def flush_body() -> None:
            normalized = re.sub(
                r"\s+",
                " ",
                " ".join(line for line in body_lines if line),
            ).strip()
            if normalized:
                paragraphs.append(normalized)
            body_lines.clear()

        for raw_line in raw_paragraph.split("\n"):
            line = re.sub(r"[ \t]+", " ", raw_line).strip()
            if not line:
                continue
            if " | " in line:
                heading, body = line.split(" | ", 1)
                if body.strip() and _looks_like_heading(heading):
                    flush_body()
                    paragraphs.append(heading)
                    body_lines.append(body)
                    continue
            if _looks_like_heading(line):
                flush_body()
                paragraphs.append(line)
            else:
                body_lines.append(line)
        flush_body()
    return "\n\n".join(paragraphs)


def load_document(text: str, file_path: str) -> ExtractedDocument:
    source_type = "text"
    warnings: tuple[str, ...] = ()
    if not text.strip() and file_path.strip():
        path = _validate_upload_path(file_path)
        extension = path.suffix.casefold()
        if extension == ".pdf":
            text, warnings = _extract_pdf(path)
            source_type = "pdf"
        elif extension == ".docx":
            text = _extract_docx(path)
            source_type = "docx"
        else:
            raise ContractError("unsupported_file_type", "Only PDF and DOCX uploads are supported.")
    normalized = _normalize_extracted_text(text)
    if not normalized:
        raise ContractError("unreadable_document", "No continuous readable text was found.")
    if len(normalized) > MAX_TEXT_CHARACTERS:
        raise ContractError("input_too_large", "The document exceeds the supported text limit.")

    first_line = next((line.strip() for line in normalized.splitlines() if line.strip()), "")
    title = first_line if len(first_line) <= 255 and not re.search(r"[.!?]\s", first_line) else ""
    return ExtractedDocument(normalized, source_type, title, warnings)


def _load_spacy() -> Any | None:
    try:
        import spacy

        try:
            return spacy.load("en_core_web_sm", disable=["ner"])
        except (OSError, ImportError):
            nlp = spacy.blank("en")
            nlp.add_pipe("sentencizer")
            return nlp
    except ImportError:
        return None


_NLP: Any | None = None
_NLP_INITIALIZED = False


def _fallback_sentences(text: str) -> list[tuple[str, int, int]]:
    output: list[tuple[str, int, int]] = []
    start = 0
    for match in re.finditer(r"[.!?](?:[\"')\]]*)\s+", text):
        punctuation_start = match.start()
        prefix = text[start:punctuation_start + 1]
        if _ABBREVIATIONS.search(prefix):
            continue
        end = match.end()
        while start < len(text) and text[start].isspace():
            start += 1
        value = text[start:punctuation_start + 1].strip()
        if value:
            output.append((value, start, punctuation_start + 1))
        start = end
    while start < len(text) and text[start].isspace():
        start += 1
    tail = text[start:].strip()
    if tail:
        output.append((tail, start, len(text)))
    return output


def _merge_place_abbreviation_boundaries(
    text: str,
    spans: list[tuple[str, int, int]],
) -> list[tuple[str, int, int]]:
    merged: list[tuple[str, int, int]] = []
    for sentence, start, end in spans:
        if merged:
            previous_sentence, previous_start, previous_end = merged[-1]
            continuation = text[start:end].strip()
            if (
                _PLACE_ABBREVIATION_END.search(previous_sentence)
                and continuation
                and (continuation.startswith("(") or continuation[0].islower())
            ):
                merged[-1] = (
                    text[previous_start:end].strip(),
                    previous_start,
                    end,
                )
                continue
        merged.append((sentence, start, end))
    return merged


def split_sentences(text: str) -> list[SourceSentence]:
    global _NLP, _NLP_INITIALIZED
    if not _NLP_INITIALIZED:
        _NLP = _load_spacy()
        _NLP_INITIALIZED = True

    records: list[SourceSentence] = []
    section = "Document"
    offset = 0
    paragraph_index = 0
    for paragraph in text.split("\n\n"):
        if not paragraph:
            offset += 2
            continue
        is_initial_title = (
            paragraph_index == 0
            and len(paragraph) <= 255
            and len(_TOKEN.findall(paragraph)) >= 3
            and not re.search(r"[.!?]\s", paragraph)
            and not paragraph.rstrip().endswith((".", "!", "?"))
        )
        if is_initial_title:
            offset += len(paragraph) + 2
            paragraph_index += 1
            continue
        if _looks_like_heading(paragraph):
            section = paragraph.strip(" :")
            offset += len(paragraph) + 2
            paragraph_index += 1
            continue

        local_spans: list[tuple[str, int, int]] = []
        if _NLP is not None:
            try:
                doc = _NLP(paragraph)
                local_spans = [
                    (sentence.text.strip(), sentence.start_char, sentence.end_char)
                    for sentence in doc.sents
                    if sentence.text.strip()
                ]
            except Exception:
                local_spans = []
        if not local_spans:
            local_spans = _fallback_sentences(paragraph)
        local_spans = _merge_place_abbreviation_boundaries(paragraph, local_spans)

        search_from = 0
        for _, local_start, local_end in local_spans:
            start = offset + local_start
            end = offset + local_end
            exact = text[start:end].strip()
            if not exact:
                continue
            # Keep a source-exact slice even if a tokenizer adjusts edge whitespace.
            found = text.find(exact, offset + search_from, offset + len(paragraph))
            if found < 0:
                continue
            start = found
            end = found + len(exact)
            search_from = end - offset
            if len(_TOKEN.findall(exact)) < 4:
                continue
            records.append(SourceSentence(
                index=len(records),
                text=exact,
                start=start,
                end=end,
                paragraph_index=paragraph_index,
                section=section,
            ))
        paragraph_index += 1
        offset += len(paragraph) + 2
    return records


def starts_with_context_reference(text: str) -> bool:
    return bool(_SENTENCE_START_REFERENCE.match(text))
