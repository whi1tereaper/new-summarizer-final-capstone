"""Additive routing, source preservation, formats and failure isolation."""
from __future__ import annotations

import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from article_detector import detect_article, threshold
from article_engine import preprocessing
from article_engine.summarizer import summarize
from evaluation.article_routing_fixtures import DEVELOPMENT, heldout
from failure_contract import ProcessingFailure
from summarizer_core.models import SourceDocument
from summary_engine_router import enabled, route_summary

ARTICLE = DEVELOPMENT[0]["text"]


def request(**kwargs):
    return {"text": ARTICLE, "automatic_article_routing": True, "summary_depth": "balanced",
            "sentence_count": 8, "analysis_mode": "general", "output_format": "paragraph", **kwargs}


class ArticleRoutingTests(unittest.TestCase):
    def setUp(self):
        # Selection itself is independent of optional POS/NER provisioning.
        self.model = patch("summarizer_core.tokenization.load_optional_spacy_pipeline", return_value=None)
        self.model.start()
        self.addCleanup(self.model.stop)

    def test_development_article_categories_have_calibrated_decisions(self):
        self.assertLessEqual(threshold(), 1)
        for d in DEVELOPMENT:
            with self.subTest(d["id"]):
                self.assertEqual(detect_article(d["text"]).is_article, d["article"])

    def test_heldout_non_articles_never_route_to_article_engine(self):
        for d in heldout():
            if not d["article"]:
                with self.subTest(d["id"]):
                    self.assertFalse(detect_article(d["text"]).is_article)

    def test_title_and_repeated_keyword_are_insufficient(self):
        self.assertFalse(detect_article("Research News Article\n\n" + "research " * 500).is_article)

    def test_missing_calibration_fails_closed(self):
        with patch("article_detector.CALIBRATION", Path("nonexistent-article-calibration.json")):
            threshold.cache_clear()
            self.assertFalse(detect_article(ARTICLE).is_article)
        threshold.cache_clear()

    def test_disabled_router_preserves_request_and_skips_extraction(self):
        original = Mock(return_value={"sentences": ["unchanged"]})
        payload = request(automatic_article_routing=False)
        with patch("summarizer_core.extraction.load_source_document") as extract:
            self.assertEqual(route_summary(payload, original), {"sentences": ["unchanged"]})
        extract.assert_not_called()
        original.assert_called_once_with(payload, None)

    def test_server_false_overrides_environment_true(self):
        with patch.dict("os.environ", {"SUMMARIZER_AUTO_ARTICLE_ROUTING": "true"}):
            self.assertFalse(enabled({"automatic_article_routing": False}))

    def test_article_success_does_not_call_original_engine(self):
        original = Mock()
        result = route_summary(request(), original)
        original.assert_not_called()
        self.assertEqual(result["source_metadata"]["engine_routing"]["selected_engine"], "recipe_article")
        self.assertFalse(result["fallback_used"])

    def test_uncertain_input_uses_original_engine_once(self):
        original = Mock(return_value={"plain_summary": "original"})
        result = route_summary(request(text="A short statement."), original)
        self.assertEqual(result["plain_summary"], "original")
        original.assert_called_once()
        self.assertFalse(result["source_metadata"]["engine_routing"]["fallback_used"])

    def test_article_failure_falls_back_once_and_records_actual_engine(self):
        original = Mock(return_value={"plain_summary": "original"})
        with patch("article_engine.summarizer.summarize", side_effect=RuntimeError("secret document path")):
            result = route_summary(request(), original)
        original.assert_called_once()
        self.assertTrue(result["fallback_used"])
        routing = result["source_metadata"]["engine_routing"]
        self.assertEqual(routing["selected_engine"], "legacy")
        self.assertEqual(routing["failure_type"], "RuntimeError")
        self.assertNotIn("secret", json.dumps(routing))

    def test_missing_article_module_keeps_original_available(self):
        original = Mock(return_value={"plain_summary": "original"})
        with patch.dict(sys.modules, {"article_engine.summarizer": None}):
            result = route_summary(request(), original)
        self.assertTrue(result["fallback_used"])
        original.assert_called_once()

    def test_invalid_article_contract_uses_original_once(self):
        original = Mock(return_value={"plain_summary": "original"})
        with patch("article_engine.summarizer.summarize", return_value={"bad": "contract"}):
            result = route_summary(request(), original)
        original.assert_called_once()
        self.assertTrue(result["fallback_used"])

    def test_extraction_error_is_not_mislabeled_as_engine_fallback(self):
        original = Mock()
        with patch("summarizer_core.extraction.load_source_document", side_effect=ProcessingFailure("FILE_CORRUPTED")):
            with self.assertRaises(ProcessingFailure):
                route_summary(request(), original)
        original.assert_not_called()

    def test_existing_synthesis_request_keeps_original_route(self):
        original = Mock(return_value={"plain_summary": "original"})
        result = route_summary(request(use_llm_synthesis=True), original)
        self.assertEqual(result["source_metadata"]["engine_routing"]["reason"], "existing_synthesis_request")
        original.assert_called_once()

    def test_fallback_reuses_extraction_without_changing_source_type(self):
        from local_cli import _summarize_request
        source = SourceDocument(ARTICLE, "docx", "Title")
        with patch("summarizer_core.extraction.load_source_document", return_value=source) as extract:
            with patch("article_engine.summarizer.summarize", side_effect=RuntimeError):
                with patch("summarizer_core.pipeline.load_source_document", side_effect=AssertionError("duplicate extraction")):
                    result = route_summary(request(), _summarize_request)
        extract.assert_called_once()
        self.assertEqual(result["source_metadata"]["source_type"], "docx")

    def test_source_spans_preserve_facts_and_qualifiers(self):
        source = "Researchers did not establish a causal effect in 2025. The sample of 1,186 people may not represent all residents."
        spans, _ = preprocessing.sentences(source)
        self.assertEqual(len(spans), 2)
        for s in spans:
            self.assertEqual(source[s.start:s.end], s.text)
        self.assertIn("1,186", spans[1].text)
        self.assertIn("may not", spans[1].text)

    def test_missing_punkt_has_source_exact_boundary_fallback(self):
        with patch("article_engine.preprocessing._tokenizer", return_value=(None, "boundary_fallback")):
            spans, status = preprocessing.sentences("Dr. Rivera measured 2.5 percent growth. The result was not significant.")
        self.assertEqual(status, "boundary_fallback")
        self.assertEqual(len(spans), 2)
        self.assertIn("Dr. Rivera", spans[0].text)

    def test_all_modes_depths_and_formats_preserve_source_evidence(self):
        source = SourceDocument(ARTICLE, "text")
        from article_engine.ranking import MODES
        for mode in MODES:
            result = summarize(request(analysis_mode=mode), source, detect_article(ARTICLE))
            self.assertEqual(result["analysis_mode"], mode)
        counts = []
        for depth in ["very_brief", "brief", "short", "balanced", "detailed", "comprehensive"]:
            values = []
            for fmt in ["paragraph", "bullets", "hybrid", "structured", "key_takeaways"]:
                result = summarize(request(summary_depth=depth, output_format=fmt), source, detect_article(ARTICLE))
                values.append(result["sentences"])
                for s, evidence in zip(result["sentences"], result["evidence"]):
                    self.assertEqual(ARTICLE[evidence["start"]:evidence["end"]], s)
            self.assertTrue(all(v == values[0] for v in values))
            counts.append(len(values[0]))
        self.assertEqual(counts, sorted(counts))

    def test_focus_changes_evidence_and_unsupported_focus_is_explicit(self):
        source = SourceDocument(ARTICLE, "text")
        focused = summarize(request(focus_query="funding operating costs"), source, detect_article(ARTICLE))
        self.assertTrue(focused["sentences"])
        unsupported = summarize(request(focus_query="volcanic seismology"), source, detect_article(ARTICLE))
        self.assertEqual(unsupported["sentences"], [])
        self.assertEqual(unsupported["source_metadata"]["article_analysis"]["focus_status"], "unsupported")
        self.assertIn("does not contain enough information", unsupported["plain_summary"])
        self.assertEqual(unsupported["evidence"], [])
        partial = summarize(request(focus_query="funding volcanic seismology"), source, detect_article(ARTICLE))
        self.assertEqual(partial["source_metadata"]["article_analysis"]["focus_status"], "unsupported")

    def test_optional_ner_unavailable_is_reported(self):
        result = summarize(request(), SourceDocument(ARTICLE, "text"), detect_article(ARTICLE))
        self.assertEqual(result["source_metadata"]["article_analysis"]["annotation_status"], "unavailable_local_model")

    def test_repeated_request_is_deterministic(self):
        a = route_summary(request(), Mock())
        b = route_summary(request(), Mock())
        self.assertEqual(a, b)

    def test_existing_engine_remains_independently_callable(self):
        from local_cli import _summarize_request
        result = _summarize_request(request(automatic_article_routing=False))
        self.assertTrue(result["sentences"])
        self.assertNotIn("engine_routing", result["source_metadata"])

    def test_nonarticle_output_matches_original_baseline(self):
        from local_cli import _summarize_request
        text = next(d["text"] for d in heldout() if d["id"] == "HO-LONG-WATER")
        payload = request(text=text)
        original = _summarize_request(payload)
        routed = route_summary(payload, _summarize_request)
        for key in ["sentences", "plain_summary", "readability", "analysis_mode", "output_format"]:
            self.assertEqual(routed[key], original[key], key)
        self.assertEqual(routed["source_metadata"]["engine_routing"]["selected_engine"], "legacy")

    def test_real_warm_supervisor_honors_per_request_rollback(self):
        from worker_daemon import WorkerSupervisor
        worker = WorkerSupervisor(90)
        try:
            status, article = worker.execute("summarize", request(), 90)
            self.assertEqual(status, 200, article)
            self.assertEqual(article["source_metadata"]["engine_routing"]["selected_engine"], "recipe_article")
            status, original = worker.execute("summarize", request(automatic_article_routing=False), 90)
            self.assertEqual(status, 200, original)
            self.assertNotIn("engine_routing", original["source_metadata"])
            self.assertFalse(original["source_metadata"]["worker_cold_start"])
        finally:
            worker.close()

    def test_docx_extraction_routes_research_article(self):
        from docx import Document
        text = DEVELOPMENT[1]["text"]
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "article.docx"
            doc = Document()
            for line in text.splitlines():
                if line:
                    doc.add_paragraph(line)
            doc.save(path)
            result = route_summary(request(text="", file_path=str(path)), Mock())
            cv = Document()
            for line in "Curriculum Vitae\nJamie Example\nEducation: Computer Science\nExperience: Research assistant\nSkills: Python and PHP".splitlines():
                cv.add_paragraph(line)
            cv_path = Path(folder) / "resume.docx"
            cv.save(cv_path)
            original = Mock(return_value={"plain_summary": "Original result"})
            cv_result = route_summary(request(text="", file_path=str(cv_path)), original)
            self.assertEqual(cv_result["source_metadata"]["engine_routing"]["selected_engine"], "legacy")
            original.assert_called_once()
            self.assertEqual(original.call_args.args[1].source_type, "docx")
        self.assertEqual(result["source_metadata"]["engine_routing"]["selected_engine"], "recipe_article")
        self.assertEqual(result["source_metadata"]["source_type"], "docx")

    def test_pdf_article_and_receipt_are_detected_from_content(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        # Controlled text PDF fixture, not a real-world extraction-quality claim.
        with tempfile.TemporaryDirectory() as folder:
            for label, text, expected in [("news", ARTICLE, True), ("receipt", "Receipt\nBread 50\nTotal 50\nCash 100\nChange 50", False)]:
                writer = PdfWriter()
                page = writer.add_blank_page(width=2000, height=2000)
                font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
                page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
                commands = ['BT /F1 12 Tf 20 1900 Td']
                for line in text.splitlines():
                    escaped = line.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
                    commands.append('(' + escaped + ') Tj 0 -18 Td')
                commands.append('ET')
                stream = DecodedStreamObject()
                stream.set_data('\n'.join(commands).encode('latin-1'))
                page[NameObject('/Contents')] = writer._add_object(stream)
                path = Path(folder) / (label + '.pdf')
                writer.write(path)
                from summarizer_core.extraction import load_source_document
                try:
                    source = load_source_document(file_path=str(path))
                except ProcessingFailure:
                    self.assertFalse(expected, "Article PDF extraction must succeed")
                    continue  # Original readability controls can reject a receipt.
                self.assertEqual(detect_article(source.raw_text).is_article, expected)
                if expected:
                    result = route_summary(request(text="", file_path=str(path)), Mock())
                    self.assertEqual(result["source_metadata"]["engine_routing"]["selected_engine"], "recipe_article")

    def test_url_content_uses_same_detector_after_existing_fetch(self):
        source = SourceDocument(DEVELOPMENT[2]["text"], "url")
        # Legacy text_utils provides a dynamic compatibility alias for fetches.
        with patch("summarizer_core.extraction.fetch_url_source", return_value=source), \
                patch("summarizer_core.text_utils.fetch_url_source", return_value=source):
            result = route_summary(request(text="https://example.org/technical-article"), Mock())
        self.assertEqual(result["source_metadata"]["engine_routing"]["selected_engine"], "recipe_article")
        self.assertEqual(result["source_metadata"]["source_type"], "url")


if __name__ == "__main__":
    unittest.main()
