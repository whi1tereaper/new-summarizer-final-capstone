"""Failure contract checks with generated documents and controlled HTTP sources."""
from __future__ import annotations

import io
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from unittest.mock import Mock, patch

ENGINE_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ENGINE_ROOT))
# Spawned fixture workers share this test's bounded numerical thread budget.
os.environ["OPENBLAS_NUM_THREADS"] = "1"
os.environ["OMP_NUM_THREADS"] = "1"

from failure_contract import ProcessingFailure, failure_response
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from summarizer_core import extraction
from light_engine.contract import ContractError, SummaryRequest, MAX_TEXT_CHARACTERS
from light_engine.extraction import _extract_pdf, _extract_docx, SourceSentence
from light_engine.validation import validate_candidates


SOURCE = (
    "The city opened a new public library on Monday. "
    "The building provides reading rooms and free internet access for residents. "
    "Local officials said the library will support education throughout the district. "
    "The library opens every morning and closes at six in the evening."
)


def pdf_fixture(path: Path, *, text: str = "", encrypted: bool = False, image_only: bool = False) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 40 700 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    if image_only:
        stream = DecodedStreamObject()
        stream.set_data(b"q 100 0 0 100 0 0 cm BI /W 1 /H 1 /CS /RGB /BPC 8 ID \xff\xff\xff EI Q")
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("private-password")
    with path.open("wb") as destination:
        writer.write(destination)


class FixtureHandler(BaseHTTPRequestHandler):
    calls: list[str] = []

    def log_message(self, *args):
        pass

    def do_GET(self):
        self.calls.append(self.path)
        status = int(self.path[1:]) if self.path[1:].isdigit() else 200
        body = {
            "/article": f"<html><head><title>Library opens</title></head><body><article><p>{SOURCE}</p></article></body></html>",
            "/short": "<html><body><article><p>The library opens on Monday morning.</p></article></body></html>",
            "/empty": "<html><body><script>renderArticle()</script></body></html>",
            "/login": '<html><body><form><input type="password"></form></body></html>',
            "/wrong.pdf": "<html><body><p>The requested document is unavailable.</p></body></html>",
        }.get(self.path, "<html><body>Request could not be completed.</body></html>").encode()
        self.send_response(status)
        self.send_header("Content-Type", "text/html")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


class FailureFeedbackTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="light-failures-")
        cls.directory = Path(cls.temp.name)
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), FixtureHandler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.base = f"http://127.0.0.1:{cls.server.server_port}"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()
        cls.temp.cleanup()

    def assert_failure(self, code, operation, *args, **kwargs):
        with self.assertRaises(ProcessingFailure) as raised:
            operation(*args, **kwargs)
        response = failure_response(raised.exception)
        self.assertEqual(response["failure"]["code"], code)
        self.assertFalse(response["ok"])
        self.assertTrue(response["failure"]["recovery"])
        self.assertEqual(len(response["failure"]["reference"]), 32)
        self.assertNotIn("summary", response)
        return response

    def test_unknown_exception_never_exposes_internal_message(self):
        response = failure_response(RuntimeError("secret document C:/private/document.pdf password=hidden"))
        self.assertEqual(response["failure"]["code"], "PROCESSING_FAILED")
        for token in ("secret document", "C:/private", "password=hidden"):
            self.assertNotIn(token, json.dumps(response))

    def test_empty_legacy_request_is_rejected_before_nlp_import(self):
        from local_cli import handle_summarize
        self.assert_failure("INPUT_EMPTY", handle_summarize, {"text": " \n "})

    def test_independent_contract_preserves_empty_and_size_diagnostics(self):
        self.assert_failure("INPUT_EMPTY", SummaryRequest.parse, {"schema_version": 1, "text": ""})
        self.assert_failure("INPUT_TOO_LARGE", SummaryRequest.parse, {"schema_version": 1, "text": "a" * (MAX_TEXT_CHARACTERS + 1)})

    def test_corrupted_pdf_is_not_described_as_scanned(self):
        path = self.directory / "corrupt.pdf"
        path.write_bytes(b"%PDF-1.7\ntruncated malformed data")
        for operation in (extraction.extract_pdf_source, _extract_pdf):
            with self.subTest(engine=operation.__module__):
                self.assert_failure("FILE_CORRUPTED", operation, path)

    def test_encrypted_pdf_never_invokes_ocr(self):
        path = self.directory / "encrypted.pdf"
        pdf_fixture(path, encrypted=True)
        with patch.object(extraction, "extract_pdf_source_with_ocr") as ocr:
            self.assert_failure("FILE_ENCRYPTED", extraction.extract_pdf_source, path)
            ocr.assert_not_called()
        self.assert_failure("FILE_ENCRYPTED", _extract_pdf, path)

    def test_legacy_pdf_without_required_user_password_stays_supported(self):
        path = self.directory / "owner-protected.pdf"
        pdf_fixture(path, text=SOURCE)
        reader = PdfReader(path)
        writer = PdfWriter()
        writer.append_pages_from_reader(reader)
        writer.encrypt(user_password="", owner_password="owner-private")
        encrypted = self.directory / "password-free.pdf"
        with encrypted.open("wb") as destination:
            writer.write(destination)
        self.assertIn("public library", extraction.extract_pdf_source(encrypted).raw_text)

    def test_blank_and_image_only_pdf_have_honest_text_diagnostic(self):
        for image_only in (False, True):
            path = self.directory / f"no-text-{image_only}.pdf"
            pdf_fixture(path, image_only=image_only)
            with self.subTest(image_only=image_only):
                response = self.assert_failure("FILE_TEXT_UNAVAILABLE", extraction.extract_pdf_source_from_reader, PdfReader(path))
                self.assertNotIn("is scanned", response["error"])
                self.assert_failure("FILE_TEXT_UNAVAILABLE", _extract_pdf, path)

    def test_supported_ocr_fallback_receives_original_failure(self):
        path = self.directory / "ocr-source.pdf"
        pdf_fixture(path, image_only=True)
        sentinel = object()
        with patch.object(extraction, "extract_pdf_source_with_ocr", return_value=sentinel) as ocr:
            self.assertIs(extraction.extract_pdf_source(path), sentinel)
            self.assertEqual(ocr.call_args.kwargs["original_failure"].failure_code, "FILE_TEXT_UNAVAILABLE")

    def test_pdf_library_failure_is_not_assumed_to_be_damage(self):
        page = Mock()
        page.extract_text.side_effect = RuntimeError("private library details")
        reader = Mock(is_encrypted=False, pages=[page])
        self.assert_failure("DOCUMENT_EXTRACTION_FAILED", extraction.extract_pdf_source_from_reader, reader)

    def test_valid_pdf_extracts_in_both_engines(self):
        path = self.directory / "valid.pdf"
        pdf_fixture(path, text=SOURCE)
        self.assertIn("public library", extraction.extract_pdf_source(path).raw_text)
        self.assertIn("public library", _extract_pdf(path)[0])

    def test_valid_empty_and_corrupt_docx(self):
        from docx import Document
        path = self.directory / "valid.docx"
        document = Document()
        document.add_paragraph(SOURCE)
        document.save(path)
        self.assertIn("public library", extraction.load_source_document(file_path=str(path)).raw_text)
        self.assertIn("public library", _extract_docx(path))
        empty = self.directory / "empty.docx"
        Document().save(empty)
        self.assert_failure("FILE_TEXT_UNAVAILABLE", extraction.load_source_document, file_path=str(empty))
        self.assert_failure("FILE_TEXT_UNAVAILABLE", _extract_docx, empty)
        corrupt = self.directory / "corrupt.docx"
        corrupt.write_bytes(b"PK\x03\x04truncated")
        self.assert_failure("FILE_CORRUPTED", extraction.load_source_document, file_path=str(corrupt))
        self.assert_failure("FILE_CORRUPTED", _extract_docx, corrupt)

    def test_invalid_unsafe_and_unreachable_urls_stay_distinct(self):
        for url in ("not a URL", "https://example.org:bad/", "https://[invalid/", "https://example.org invalid/"):
            self.assert_failure("URL_INVALID", extraction.validate_public_url_for_fetch, url)
        for url in ("http://localhost/private", "http://127.0.0.1/private", "http://[::1]/", "https://user:password@example.org/"):
            self.assert_failure("URL_UNSAFE", extraction.validate_public_url_for_fetch, url)
        with patch.object(socket, "getaddrinfo", side_effect=socket.gaierror()):
            self.assert_failure("URL_UNREACHABLE", extraction.validate_public_url_for_fetch, "https://source.invalid/")

    def test_http_restrictions_and_missing_content(self):
        for path, code in (("403", "URL_ACCESS_RESTRICTED"), ("401", "URL_ACCESS_RESTRICTED"), ("404", "URL_CONTENT_UNAVAILABLE"), ("429", "URL_RATE_LIMITED"), ("503", "URL_UNREACHABLE"), ("empty", "URL_CONTENT_UNAVAILABLE"), ("login", "URL_ACCESS_RESTRICTED"), ("wrong.pdf", "URL_CONTENT_UNAVAILABLE")):
            with self.subTest(path=path), patch.object(extraction, "validate_public_url_for_fetch"):
                self.assert_failure(code, extraction.fetch_url_source, self.base + "/" + path)

    def test_valid_and_short_articles_fetch_once(self):
        for path in ("article", "short"):
            FixtureHandler.calls.clear()
            with patch.object(extraction, "validate_public_url_for_fetch"):
                document = extraction.fetch_url_source(self.base + "/" + path)
            self.assertIn("library", document.raw_text)
            self.assertEqual(FixtureHandler.calls, ["/" + path])

    def test_network_timeout_is_not_reported_as_invalid_url(self):
        import requests
        with patch.object(extraction, "validate_public_url_for_fetch"), patch.object(requests, "get", side_effect=requests.exceptions.Timeout()):
            self.assert_failure("URL_TIMEOUT", extraction.fetch_public_url_response, "https://example.org/")

    def test_quality_validation_has_distinct_failure(self):
        source = "The library opens on Monday morning."
        candidate = SourceSentence(0, "The library opens on Tuesday morning.", 0, len(source), 0, "Document")
        self.assert_failure("SUMMARY_QUALITY_FAILED", validate_candidates, [candidate], source)

    def test_cli_errors_are_json_without_exception_details(self):
        for script, payload, expected in (("local_cli.py", {"text": ""}, "INPUT_EMPTY"), ("light_engine_cli.py", {"schema_version": 1, "text": ""}, "INPUT_EMPTY")):
            result = subprocess.run([sys.executable, "-B", str(ENGINE_ROOT / script), "summarize"], input=json.dumps(payload), text=True, capture_output=True, timeout=30)
            self.assertEqual(result.returncode, 2)
            response = json.loads(result.stderr)
            self.assertEqual(response["failure"]["code"], expected)
            self.assertEqual(result.stdout, "")

    def test_daemon_preserves_document_error_and_reuses_live_child(self):
        from worker_daemon import WorkerSupervisor
        path = self.directory / "worker-corrupt.pdf"
        path.write_bytes(b"%PDF-1.7\ntruncated")
        worker = WorkerSupervisor(30)
        try:
            status, response = worker.execute("extract-pdf", {"file_path": str(path)}, 30)
            self.assertEqual(status, 422)
            self.assertEqual(response["failure"]["code"], "FILE_CORRUPTED")
            worker_pid = worker.process.pid
            self.assertEqual(worker.execute("self-test", {}, 30)[0], 200)
            self.assertEqual(worker.process.pid, worker_pid)
        finally:
            worker.close()


class PipelineFailureTests(unittest.TestCase):
    @classmethod
    def tearDownClass(cls):
        from summarizer_core import grammar_check
        if grammar_check._GLOBAL_CHECKER is not None:
            grammar_check._GLOBAL_CHECKER.close()
            grammar_check._GLOBAL_CHECKER = None

    def test_short_valid_text_retains_supported_behavior(self):
        from summarizer_core.models import SummarizationRequest
        from summarizer_core.pipeline import SummarizationPipeline
        source = "The public library opens on Monday morning and provides free internet access for residents."
        result = SummarizationPipeline().summarize(SummarizationRequest(text=source))
        self.assertIn("library", result.plain_summary)

    def test_legacy_minimum_is_classified_without_changing_quality_gate(self):
        from summarizer_core.models import SummarizationRequest
        from summarizer_core.pipeline import SummarizationPipeline
        with self.assertRaises(ProcessingFailure) as raised:
            SummarizationPipeline().summarize(SummarizationRequest(text="The public library opens on Monday morning."))
        self.assertEqual(raised.exception.failure_code, "INPUT_TOO_SHORT")

    def test_no_candidates_and_empty_quality_result_are_distinct(self):
        from summarizer_core.models import SummarizationRequest
        from summarizer_core.pipeline import SummarizationPipeline
        from summarizer_core.fact_ledger import FactualityValidator
        request = SummarizationRequest(text=SOURCE)
        with patch.object(SummarizationPipeline, "_build_sentence_candidates", return_value=([], [])):
            with self.assertRaises(ProcessingFailure) as raised:
                SummarizationPipeline().summarize(request)
            self.assertEqual(raised.exception.failure_code, "NLP_NO_VALID_CANDIDATES")
        with patch.object(FactualityValidator, "validate_and_repair", return_value=([], {"status": "rejected"})):
            with self.assertRaises(ProcessingFailure) as raised:
                SummarizationPipeline().summarize(request)
            self.assertEqual(raised.exception.failure_code, "SUMMARY_QUALITY_FAILED")


if __name__ == "__main__":
    unittest.main()
