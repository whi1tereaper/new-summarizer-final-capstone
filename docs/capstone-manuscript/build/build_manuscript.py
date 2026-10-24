"""Assemble the evidence-grounded Markdown manuscript as an editable Word file.

This script only reads source material and writes files beneath docs/.  See
README.md for the source conventions and command line.  Requires python-docx.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import sys
from urllib.parse import unquote

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parents[1]
DOCS = BASE.parent.resolve()
DEFAULT_FILES = [
    "PRELIMINARY_PAGES.md", "CHAPTER_1.md", "CHAPTER_2.md", "CHAPTER_3.md",
    "CHAPTER_4.md", "REFERENCES.md", "APPENDICES_A_E.md",
    "APPENDIX_F_USER_MANUAL.md", "APPENDIX_G_MAINTENANCE_MANUAL.md",
    "EVALUATION_INSTRUMENTS.md", "ADDITIONAL_APPENDICES.md",
]
MANDATORY_FILES = set(DEFAULT_FILES) - {"APPENDICES_A_E.md", "ADDITIONAL_APPENDICES.md"}
PAGE_WIDTH = 515.88
PAGE_HEIGHT = 728.52
TEXT_WIDTH = PAGE_WIDTH - 108 - 72
TEXT_HEIGHT = PAGE_HEIGHT - 144
CAPTION_RE = re.compile(r"^(Figure|Table)\s+(\d+)[.:]\s+(.*)$", re.I)
IMAGE_RE = re.compile(r'^!\[([^\]]*)\]\((.+?)\)(?:\s*)$')
INLINE_RE = re.compile(r"(!?\[[^\]]*\]\([^\)]+\)|`[^`]+`|\*\*[^*]+\*\*|__[^_]+__|(?<!\*)\*[^*]+\*(?!\*))")


def local_output(path: Path) -> Path:
    path = path.resolve()
    if not path.is_relative_to(DOCS):
        raise ValueError(f"Output must remain beneath docs/: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def xml(tag, **attrs):
    element = OxmlElement(tag)
    for name, value in attrs.items():
        element.set(qn(name), str(value))
    return element


def set_font(run, size=12, bold=False, italic=False):
    run.font.name = "Courier New"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
    if fonts is not None:
        for name in ("ascii", "hAnsi", "cs", "eastAsia"):
            fonts.set(qn(f"w:{name}"), "Courier New")
    return run


def field(paragraph, instruction, cached="", size=12, bold=False):
    run = paragraph.add_run()
    set_font(run, size, bold)
    run._r.append(xml("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
    text = xml("w:instrText", **{"xml:space": "preserve"})
    text.text = f" {instruction} "
    run._r.append(text)
    run._r.append(xml("w:fldChar", **{"w:fldCharType": "separate"}))
    value = xml("w:t", **{"xml:space": "preserve"})
    value.text = cached
    run._r.append(value)
    run._r.append(xml("w:fldChar", **{"w:fldCharType": "end"}))
    return run


def inline(paragraph, value, size=12, bold=False):
    value = value.replace("<br/>", "\n").replace("<br />", "\n").replace("<br>", "\n")
    for token in INLINE_RE.split(value):
        if not token:
            continue
        if token.startswith("[") and "](" in token:
            match = re.match(r"\[([^\]]*)\]\(([^)]+)\)", token)
            label, href = match.groups()
            hyperlink = xml("w:hyperlink")
            relation = paragraph.part.relate_to(href, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            hyperlink.set(qn("r:id"), relation)
            run = paragraph.add_run(label)
            set_font(run, size, bold)
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
        elif token.startswith("**") or token.startswith("__"):
            set_font(paragraph.add_run(token[2:-2]), size, True)
        elif token.startswith("`"):
            set_font(paragraph.add_run(token[1:-1]), size, bold)
        elif token.startswith("*"):
            set_font(paragraph.add_run(token[1:-1]), size, bold, True)
        else:
            set_font(paragraph.add_run(token), size, bold)


def configure_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Courier New"
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = Pt(20.4)
    normal.paragraph_format.first_line_indent = Inches(0.25)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for n in range(1, 5):
        style = doc.styles[f"Heading {n}"]
        style.font.name = "Courier New"
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        pf = style.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.line_spacing = Pt(20.4)
        pf.space_before = Pt(12 if n > 1 else 0)
        pf.space_after = Pt(12)
        pf.keep_with_next = True
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if n == 1 else WD_ALIGN_PARAGRAPH.LEFT
    for name in ("Caption", "List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Courier New"
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.line_spacing = Pt(16)
        style.paragraph_format.space_after = Pt(6)
    for name in ("TOC 1", "TOC 2", "TOC 3"):
        if name not in doc.styles:
            doc.styles.add_style(name, 1)
        style = doc.styles[name]
        style.font.name = "Courier New"
        style.font.size = Pt(10)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.line_spacing = Pt(14)
        style.paragraph_format.space_after = Pt(3)
    settings = doc.settings.element
    settings.append(xml("w:updateFields", **{"w:val": "true"}))
    configure_section(doc.sections[0], "lowerRoman", first_page=True)
    doc.core_properties.title = "LIGHT summarizer — Evidence-grounded Capstone Manuscript Draft"
    doc.core_properties.subject = "Editable academic manuscript with pending research information clearly identified"
    doc.core_properties.author = "Robert Martin Baltisoto; Shielo Mariell Leodones; Gieselle Sabbaluca; Sofiah Heaven Galman"
    doc.core_properties.keywords = "audit; capstone; draft; evidence"


def configure_section(section, number_format="decimal", first_page=False):
    section.page_width = Pt(PAGE_WIDTH)
    section.page_height = Pt(PAGE_HEIGHT)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.5)
    section.different_first_page_header_footer = first_page
    section.footer.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False
    pg_num = section._sectPr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = xml("w:pgNumType")
        section._sectPr.append(pg_num)
    pg_num.set(qn("w:fmt"), number_format)
    pg_num.set(qn("w:start"), "1")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.first_line_indent = Pt(0)
    field(footer, "PAGE", "1")


class Builder:
    def __init__(self, files, output):
        self.files = files
        self.output = output
        self.doc = Document()
        configure_document(self.doc)
        self.report = {"source_files": [], "source_sha256": {}, "warnings": [], "figures": [], "tables": [], "images": [], "headings": []}
        self.counters = {"Figure": 0, "Table": 0}
        self.front = True
        self.main_started = False
        self.has_content = False
        self.current_heading = ""
        self.indexes = set()
        self.last_caption = None
        self.skip_caption = False
        self.centered = False

    def paragraph(self, text, *, style=None, small=False):
        p = self.doc.add_paragraph(style=style)
        inline(p, text, 9 if small else 12)
        if small:
            p.paragraph_format.line_spacing = Pt(12)
            p.paragraph_format.first_line_indent = Pt(0)
        if self.front:
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = Pt(16)
            p.paragraph_format.space_after = Pt(4)
        if self.centered:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
        self.has_content = True
        self.last_caption = None
        return p

    def heading(self, title, level):
        title = re.sub(r"\*\*([^*]+)\*\*", r"\1", title)
        if self.front and title.lower() == "abstract":
            for kind in ("TOC", "LIST_OF_TABLES", "LIST_OF_FIGURES"):
                if kind not in self.indexes:
                    self.index(kind, include_heading=True)
        new_page = level == 1 or (self.front and level == 2)
        p = self.doc.add_paragraph(style=f"Heading {min(level, 4)}")
        if self.front:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if new_page and self.has_content:
            p.paragraph_format.page_break_before = True
        if level == 1:
            title = title.upper()
            match = re.match(r"^(CHAPTER\s+[IVXLCDM\d]+)\s*[—–:\-]?\s*(.*)$", title)
            if match and match.group(2):
                title = match.group(1) + "\n" + match.group(2)
        inline(p, title, bold=True)
        self.has_content = True
        self.current_heading = title.replace("\n", " — ")
        self.report["headings"].append({"level": level, "title": self.current_heading})
        self.last_caption = None
        return p

    def index(self, kind, include_heading=False):
        if kind in self.indexes:
            return
        titles = {"TOC": "Table of Contents", "LIST_OF_TABLES": "List of Tables", "LIST_OF_FIGURES": "List of Figures"}
        commands = {"TOC": 'TOC \\o "1-3" \\h \\z \\u', "LIST_OF_TABLES": 'TOC \\h \\z \\c "Table"', "LIST_OF_FIGURES": 'TOC \\h \\z \\c "Figure"'}
        if include_heading:
            self.heading(titles[kind], 1)
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = Pt(14)
        field(p, commands[kind], "Update this field in Word, or rebuild with --pdf, to display entries and page numbers.", 10)
        self.indexes.add(kind)
        self.has_content = True

    def caption(self, kind, number, title, generated=False):
        previous = self.counters[kind]
        if number != previous + 1:
            self.report["warnings"].append(f"{kind} numbering jumps from {previous} to {number}: {title}")
        self.counters[kind] = number
        p = self.doc.add_paragraph(style="Caption")
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind == "Figure" else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = kind == "Table"
        p.paragraph_format.keep_together = True
        inline(p, f"{kind} ", bold=True)
        field(p, f"SEQ {kind} \\r {number} \\* ARABIC", str(number), bold=True)
        inline(p, f". {title}", bold=True)
        self.report["figures" if kind == "Figure" else "tables"].append({"number": number, "title": title, "generated_caption": generated})
        self.last_caption = kind
        self.has_content = True

    def image(self, alt, path, source, next_line=""):
        clean_path = re.sub(r'\s+["\'].*["\']$', "", path.strip()).strip("<>")
        if re.match(r"https?://", clean_path):
            raise ValueError(f"Manuscript images must be local evidence: {clean_path}")
        image_path = (source.parent / unquote(clean_path)).resolve()
        if image_path.suffix.lower() == ".svg":
            png = image_path.with_suffix(".png")
            if not png.exists():
                png = local_output(BASE / "build" / "rasterized" / (image_path.stem + ".png"))
                cmd = ["node", str(BASE / "build" / "rasterize_svg.mjs"), str(image_path), str(png)]
                subprocess.run(cmd, check=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            image_path = png
        if not image_path.exists():
            raise FileNotFoundError(f"Missing manuscript image: {image_path}")
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        shape = run.add_picture(str(image_path), width=Pt(TEXT_WIDTH))
        max_height = Pt(TEXT_HEIGHT - 95)
        if shape.height > max_height:
            ratio = max_height / shape.height
            shape.width = int(shape.width * ratio)
            shape.height = max_height
        shape._inline.docPr.set("descr", alt)
        self.report["images"].append({"source": str(image_path.relative_to(DOCS)), "alt": alt})
        self.has_content = True
        self.last_caption = None
        next_caption = CAPTION_RE.match(next_line.strip().strip("*"))
        if not next_caption:
            caption = CAPTION_RE.match(alt.strip().strip("*"))
            if caption:
                self.caption(caption[1].title(), int(caption[2]), caption[3])
            else:
                self.caption("Figure", self.counters["Figure"] + 1, alt or image_path.stem, generated=True)

    def table(self, rows):
        if self.last_caption != "Table" and not self.skip_caption:
            self.caption("Table", self.counters["Table"] + 1, self.current_heading.title(), generated=True)
        self.skip_caption = False
        width = max(len(row) for row in rows)
        if width < 1:
            return
        table = self.doc.add_table(rows=len(rows), cols=width)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        # Data dictionaries retain a wider human-readable description column.
        if width == 4 and "field" in rows[0][0].lower():
            proportions = [0.23, 0.35, 0.20, 0.22]
        elif width == 3:
            proportions = [0.43, 0.285, 0.285]
        else:
            proportions = [1 / width] * width
        for column, proportion in zip(table.columns, proportions):
            column.width = Pt(TEXT_WIDTH * proportion)
        for i, (row, data) in enumerate(zip(table.rows, rows)):
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(xml("w:cantSplit"))
            if i == 0:
                tr_pr.append(xml("w:tblHeader", **{"w:val": "true"}))
            for j, cell in enumerate(row.cells):
                cell.width = Pt(TEXT_WIDTH * proportions[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if i == 0:
                    cell._tc.get_or_add_tcPr().append(xml("w:shd", **{"w:fill": "E6E6E6"}))
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = Pt(11.5)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                inline(p, data[j] if j < len(data) else "", 8 if width > 4 else 9, i == 0)
        self.has_content = True
        self.last_caption = None

    def parse(self, source, included=False):
        lines = source.read_text(encoding="utf-8-sig").splitlines()
        self.report["source_files"].append(os.path.relpath(source, BASE))
        self.report["source_sha256"][os.path.relpath(source, BASE)] = hashlib.sha256(source.read_bytes()).hexdigest()
        if source.name != "PRELIMINARY_PAGES.md" and not self.main_started and not included:
            section = self.doc.add_section(WD_SECTION_START.NEW_PAGE)
            configure_section(section)
            self.front = False
            self.main_started = True
            self.has_content = False
        i = 0
        code = False
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue
            if line.startswith("```"):
                code = not code
                i += 1
                continue
            if code:
                self.paragraph(lines[i], small=True)
                i += 1
                continue
            if line.startswith("<!--"):
                marker = line.removeprefix("<!--").removesuffix("-->").strip()
                if marker in ("TOC", "LIST_OF_TABLES", "LIST_OF_FIGURES"):
                    self.index(marker)
                elif marker == "PAGEBREAK":
                    self.doc.add_page_break()
                elif marker == "NO_TABLE_CAPTION":
                    self.skip_caption = True
                elif marker == "CENTER":
                    self.centered = True
                elif marker == "END_CENTER":
                    self.centered = False
                elif marker.startswith("INCLUDE "):
                    include = (source.parent / marker[8:].strip()).resolve()
                    if not include.is_relative_to(DOCS):
                        raise ValueError("INCLUDE paths must be documentation files beneath docs/.")
                    self.parse(include, included=True)
                i += 1
                continue
            heading = re.match(r"^(#{1,6})\s+(.*)$", line)
            if heading:
                self.heading(heading[2], len(heading[1]))
                i += 1
                continue
            if re.match(r"^(-{3,}|\*{3,}|_{3,})$", line):
                i += 1
                continue
            image = IMAGE_RE.match(line)
            if image:
                j = i + 1
                while j < len(lines) and not lines[j].strip():
                    j += 1
                self.image(image[1], image[2], source, lines[j] if j < len(lines) else "")
                i += 1
                continue
            caption = CAPTION_RE.match(line.strip("*"))
            if caption:
                self.caption(caption[1].title(), int(caption[2]), caption[3])
                i += 1
                continue
            if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
                def cells(text):
                    return [cell.replace(r"\|", "|").strip() for cell in re.split(r"(?<!\\)\|", text.strip().strip("|"))]
                rows = [cells(line)]
                i += 2
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rows.append(cells(lines[i]))
                    i += 1
                self.table(rows)
                continue
            bullet = re.match(r"^[-*+]\s+(.*)$", line)
            numbered = re.match(r"^(\d+)[.)]\s+(.*)$", line)
            if bullet or numbered:
                # Retain explicit source numbers: automatic Word lists otherwise
                # continue numbering across unrelated procedures.
                p = self.paragraph(bullet[1] if bullet else f"{numbered[1]}. {numbered[2]}", style="List Bullet" if bullet else None)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Pt(0)
                if numbered:
                    p.paragraph_format.left_indent = Inches(0.2)
                    p.paragraph_format.first_line_indent = Inches(-0.2)
                i += 1
                continue
            paragraph = [line.lstrip("> ") if line.startswith(">") else line]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,6}\s|\||!\[|<!--|```|[-*+]\s|\d+[.)]\s|Figure\s+\d|Table\s+\d)", lines[i].strip()):
                paragraph.append(lines[i].strip())
                i += 1
            self.paragraph(" ".join(paragraph))
        if source.name == "PRELIMINARY_PAGES.md":
            for kind in ("TOC", "LIST_OF_TABLES", "LIST_OF_FIGURES"):
                if kind not in self.indexes:
                    self.index(kind, include_heading=True)

    def build(self):
        for source in self.files:
            self.parse(source)
        local_output(self.output)
        self.doc.save(self.output)
        self.report["output_docx"] = str(self.output.relative_to(DOCS))
        self.report["layout"] = {"page_width_pt": PAGE_WIDTH, "page_height_pt": PAGE_HEIGHT, "body_font": "Courier New", "body_font_pt": 12, "line_spacing_pt": 20.4, "margins_pt": {"left": 108, "right": 72, "top": 72, "bottom": 72}}
        return self.report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, help="JSON list of ordered Markdown paths relative to the manuscript folder")
    parser.add_argument("--output", type=Path, default=BASE / "EDITABLE_MANUSCRIPT.docx")
    parser.add_argument("--pdf", action="store_true", help="Update Word indexes and export PDF through local LibreOffice")
    parser.add_argument("--allow-missing", action="store_true", help="Allow incomplete sources for an explicitly partial draft")
    parser.add_argument("--libreoffice", type=Path, default=Path("C:/Program Files/LibreOffice/program/soffice.exe"))
    args = parser.parse_args()
    names = json.loads(args.manifest.read_text(encoding="utf-8-sig")) if args.manifest else DEFAULT_FILES
    if isinstance(names, dict):
        names = names["files"]
    sources = [(BASE / name).resolve() for name in names]
    missing = [path for path in sources if not path.exists()]
    required_missing = [path for path in missing if args.manifest or path.name in MANDATORY_FILES]
    if required_missing and not args.allow_missing:
        parser.error("Missing manuscript source files: " + ", ".join(path.name for path in required_missing))
    sources = [path for path in sources if path.exists()]
    if not sources:
        parser.error("No manuscript sources exist.")
    output = local_output(args.output)
    report = Builder(sources, output).build()
    report["omitted_missing_files"] = [str(path.relative_to(BASE)) for path in missing]
    if args.pdf:
        lo_python = args.libreoffice.parent / "python.exe"
        if not lo_python.exists():
            raise FileNotFoundError("LibreOffice's bundled Python is required to update all manuscript indexes.")
        cmd = [str(lo_python), str(BASE / "build" / "update_and_export.py"), str(output), "--soffice", str(args.libreoffice)]
        subprocess.run(cmd, check=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        report["output_pdf"] = str(output.with_suffix(".pdf").relative_to(DOCS))
    report_path = local_output(BASE / "build" / "BUILD_REPORT.json")
    report_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"docx": str(output), "pdf": str(output.with_suffix('.pdf')) if args.pdf else None, "sources": len(sources), "figures": len(report['figures']), "tables": len(report['tables']), "warnings": report['warnings']}, ensure_ascii=False))


if __name__ == "__main__":
    main()
