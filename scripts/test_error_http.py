"""Exercise real XAMPP HTTP upload, validation, worker and job contracts."""
from __future__ import annotations
import importlib.util, json, os, re, sys, time, uuid
from pathlib import Path
import requests
from docx import Document
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"python-engine"))
spec=importlib.util.spec_from_file_location("failure_fixtures",ROOT/"python-engine/tests/test_failure_feedback.py")
fixtures=importlib.util.module_from_spec(spec); spec.loader.exec_module(fixtures)
OUT=ROOT/"artifacts/error-feedback"; OUT.mkdir(parents=True,exist_ok=True)
SOURCE=fixtures.SOURCE
fixtures.pdf_fixture(OUT/"valid.pdf",text=SOURCE)
fixtures.pdf_fixture(OUT/"encrypted.pdf",text=SOURCE,encrypted=True)
fixtures.pdf_fixture(OUT/"image-only.pdf",image_only=True)
(OUT/"corrupted.pdf").write_bytes(b"%PDF-1.7\nbroken document\n%%EOF")
doc=Document(); doc.add_paragraph(SOURCE); doc.save(OUT/"valid.docx")
(OUT/"unsupported.txt").write_text("This is not a supported file.")
doc=Document(); doc.save(OUT/"empty.docx")
BASE=os.environ.get("LIGHT_TEST_URL", "http://localhost/app/public/")
previous=json.loads((OUT/"http-results.json").read_text()) if (OUT/"http-results.json").exists() else {}
report={"base":BASE,"cases":[],"job_ids":previous.get("job_ids",[])}
def session():
    client=requests.Session()
    page=client.get(BASE+"summarizer.php",timeout=15)
    csrf=re.search(r'name="csrf_token" value="([^"]+)"',page.text)
    assert csrf, "form csrf rendered"
    return client,csrf.group(1)
client,csrf=session()
fields={"csrf_token":csrf,"source_type":"text","original_text":SOURCE,"document_title":"LIGHT error HTTP verification","analysis_mode":"general","output_format":"paragraph","summary_depth":"brief"}
def record(name,expected,response):
    data=response.json()
    actual=data.get("failure",{}).get("code") if not data.get("success") else data.get("status","accepted")
    assert actual==expected,(name,expected,response.status_code,data)
    assert not re.search(r"[A-Z]:\\|Traceback|stack trace|worker_daemon\.py",json.dumps(data)),name+" safe public body"
    report["cases"].append({"name":name,"expected":expected,"actual":actual,"http_status":response.status_code,"verified":True})
    print("PASS",name,actual,response.status_code,flush=True)
    return data
def submit(overrides=None,file=None,token=None):
    body=fields|{"request_token":token or uuid.uuid4().hex}|(overrides or {})
    files={}
    if file:
        body["source_type"]="file"; body["original_text"]=""
        path=OUT/file
        files={"pdf_file":(file,path.read_bytes(),"application/octet-stream")}
    return client.post(BASE+"summarize.php",data=body,files=files,headers={"Accept":"application/json"},timeout=45)
def poll(job):
    deadline=time.time()+210
    while time.time()<deadline:
        response=client.get(BASE+"summarize.php",params={"action":"status","job_id":job},timeout=20)
        if response.json().get("status") not in {"queued","processing"}: return response
        time.sleep(1)
    raise AssertionError("job did not settle")
try:
    record("empty server input","INPUT_EMPTY",submit({"original_text":" "}))
    record("unreadable server input","CONTENT_UNREADABLE",submit({"original_text":"***"}))
    record("oversized server input","INPUT_TOO_LARGE",submit({"original_text":"x"*200001}))
    record("invalid server URL","URL_INVALID",submit({"source_type":"url","original_text":"","source_url":"not a URL"}))
    record("unsafe server URL","URL_UNSAFE",submit({"source_type":"url","original_text":"","source_url":"http://127.0.0.1/private"}))
    record("unsupported actual upload","FILE_UNSUPPORTED",submit(file="unsupported.txt"))
    record("csrf submission","ACCESS_DENIED",submit({"csrf_token":"wrong"}))
    record("GET execute forbidden","REQUEST_INVALID",client.get(BASE+"summarize.php?action=execute",timeout=10))
    for filename,expected in [("corrupted.pdf","FILE_CORRUPTED"),("encrypted.pdf","FILE_ENCRYPTED"),("image-only.pdf","FILE_TEXT_UNAVAILABLE"),("empty.docx","FILE_TEXT_UNAVAILABLE"),("valid.pdf","completed"),("valid.docx","completed"),(None,"completed")]:
        token=uuid.uuid4().hex
        accepted=submit(file=filename,token=token)
        data=accepted.json()
        assert data.get("success"),data
        job=data["job_id"]; report["job_ids"].append(job)
        repeated=submit(file=filename,token=token).json()
        assert repeated["job_id"]==job,"submission idempotency"
        other,_=session()
        record((filename or "text")+" foreign job access","JOB_NOT_FOUND",other.get(BASE+"summarize.php",params={"action":"status","job_id":job},timeout=10))
        record((filename or "text")+" mutation csrf","ACCESS_DENIED",client.post(BASE+"summarize.php?action=cancel",data={"job_id":job,"csrf_token":"wrong"},timeout=10))
        execution=client.post(BASE+"summarize.php?action=execute",data={"job_id":job,"csrf_token":csrf},timeout=20)
        terminal=poll(job)
        final=record(filename or "valid pasted text",expected,terminal)
        refreshed=client.get(BASE+"summarize.php",params={"action":"status","job_id":job},timeout=10).json()
        assert refreshed.get("failure",{}).get("reference")==final.get("failure",{}).get("reference"),"stable reference"
        if expected=="completed":
            assert refreshed["redirect"]==final["redirect"]
            again=submit(file=filename,token=token).json()
            assert again["job_id"]==job and again["redirect"]==final["redirect"],"completed job reused"
        else:
            assert final["can_retry"] is False,"permanent file error cannot repeat"
    # Real public URL, with full production SSRF policy in place.
    accepted=submit({"source_type":"url","original_text":"","source_url":"https://www.iana.org/help/example-domains"}).json()
    if accepted.get("success"):
        job=accepted["job_id"]; report["job_ids"].append(job)
        client.post(BASE+"summarize.php?action=execute",data={"job_id":job,"csrf_token":csrf},timeout=20)
        final=poll(job)
        if final.json().get("status")=="completed": record("valid live article URL","completed",final)
        else:
            report["cases"].append({"name":"live article URL","verified":False,"observed":final.json()})
            print("LIVE URL DID NOT COMPLETE",final.json(),flush=True)
    else:
        report["cases"].append({"name":"live article URL","verified":False,"observed":accepted})
    # Existing legacy minimum should classify accurately once its actual gate runs.
    accepted=submit({"original_text":"The library opens tomorrow."}).json()
    job=accepted["job_id"]; report["job_ids"].append(job)
    client.post(BASE+"summarize.php?action=execute",data={"job_id":job,"csrf_token":csrf},timeout=20)
    record("insufficient legacy source","INPUT_TOO_SHORT",poll(job))
finally:
    (OUT/"http-results.json").write_text(json.dumps(report,indent=2),encoding="utf-8")
